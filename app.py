from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pptx import Presentation
from pptx.util import Inches

app = FastAPI(title="PDF to PPT API", version="1.0.0")


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/convert/pdf-to-ppt")
async def convert_pdf_to_ppt(file: UploadFile = File(...)) -> FileResponse:
    if file.content_type not in {"application/pdf", "application/octet-stream"}:
        raise HTTPException(status_code=400, detail="PDF 파일만 업로드할 수 있습니다.")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="업로드된 파일이 비어 있습니다.")

    try:
        pdf_document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise HTTPException(status_code=400, detail="유효한 PDF가 아닙니다.") from exc

    if pdf_document.page_count == 0:
        raise HTTPException(status_code=400, detail="페이지가 없는 PDF입니다.")

    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    for i in range(pdf_document.page_count):
        page = pdf_document.load_page(i)
        pix = page.get_pixmap(dpi=200)
        image_bytes = pix.tobytes("png")

        slide = presentation.slides.add_slide(blank_layout)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as image_temp:
            image_temp.write(image_bytes)
            image_path = image_temp.name

        try:
            slide.shapes.add_picture(
                image_path,
                left=Inches(0),
                top=Inches(0),
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        finally:
            os.unlink(image_path)

    output_stream = io.BytesIO()
    presentation.save(output_stream)
    output_stream.seek(0)

    output_dir = Path(tempfile.mkdtemp(prefix="pdf2ppt_"))
    original_name = Path(file.filename or "converted").stem
    output_path = output_dir / f"{original_name}.pptx"
    output_path.write_bytes(output_stream.read())

    return FileResponse(
        path=output_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=output_path.name,
    )
